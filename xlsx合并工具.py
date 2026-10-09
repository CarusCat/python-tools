import os
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill

try:
    folder = input("请粘贴要合并的文件夹路径，然后按回车：\n").strip('"')

    # 表格合并
    wb_out = Workbook()
    ws_out = wb_out.active
    standard_header = None

    header_font = Font(bold=True,color="FFFFFF")
    header_fill = PatternFill("solid",fgColor="A3D5E8")

    seen = set()
    list_name = []

    for filename in sorted(os.listdir(folder)):
        if not filename.endswith(".xlsx"):
            continue
        list_name.append(filename)
        wb_in = load_workbook(os.path.join(folder,filename))
        ws_in = wb_in.active
        header = next(ws_in.iter_rows(min_row=1,max_row=1,values_only=True))
        if standard_header is None:
            standard_header = header 
            ws_out.append(list(header) + ["来源"])
        elif header != standard_header:
            print("警告",filename,"的表头",header,"与标准",standard_header,"不一致")
        for row in ws_in.iter_rows(min_row=2,values_only=True):
            name = row[0]
            if name in seen:
                print("跳过重复：",name,"来自",filename)
                continue
            seen.add(name)
            ws_out.append(list(row) + [filename])
    if standard_header is None:
        print("没找到任何Excel文件")
    else:
        for cell in ws_out[1]:
            cell.font = header_font
            cell.fill = header_fill

        wb_out.save(os.path.join(folder,"合并总表.xlsx"))
        print(f"合并完成！共处理{len(list_name)}个文件，总表{ws_out.max_row - 1}行")

except Exception as e:
    print("出错了：",e)

input("按回车退出...")
exit()